"""Prepara una importación auditada de la agenda, sin escribir en Supabase.

Requiere openpyxl y phonenumbers==9.0.41. Ejemplo:
python scripts/agenda_import.py Informacion/Todos_los_contactos.xlsx --output-dir <directorio>
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

import phonenumbers
from openpyxl import load_workbook
from phonenumbers import PhoneNumberFormat, PhoneNumberType


COUNTRIES = {
    "UY": "Uruguay", "AR": "Argentina", "BR": "Brasil", "CL": "Chile",
    "PY": "Paraguay", "ES": "España", "US": "Estados Unidos", "MX": "México",
    "CO": "Colombia", "PE": "Perú", "BO": "Bolivia", "VE": "Venezuela",
    "EC": "Ecuador", "IT": "Italia", "FR": "Francia", "DE": "Alemania",
}
CATEGORY_PATTERNS = {
    "Mates": r"\b(mate|mates|imperial|camionero|torpedo|criollo|coquito|porongo)\b",
    "Bombillas": r"\b(bombilla|bombillas|bombillon|bombillones)\b",
    "Materas": r"\b(matera|materas)\b",
    "Termos": r"\b(termo|termos|thermos)\b",
    "Yerberos": r"\b(yerbero|yerbera|yerberos|yerberas)\b",
    "Posa mates": r"\b(posa\s?mate|posa\s?mates|posamate|posamates)\b",
    "Kits materos": r"\b(kit\s+matero|kit\s+materos)\b",
    "Billeteras": r"\b(billetera|billeteras)\b",
    "Botas": r"\b(bota|botas)\b",
    "Carteras": r"\b(cartera|carteras)\b",
    "Cintos": r"\b(cinto|cintos)\b",
    "Cuchillos": r"\b(cuchillo|cuchillos)\b",
}
EMAIL = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")


def clean(value: object) -> str:
    return " ".join(str(value or "").split()).strip()


def fold(value: str) -> str:
    return "".join(ch for ch in unicodedata.normalize("NFKD", value.casefold()) if not unicodedata.combining(ch))


def category_for(name: str) -> str | None:
    folded = fold(name)
    found = {category for category, pattern in CATEGORY_PATTERNS.items() if re.search(pattern, folded)}
    if "Kits materos" in found:
        found.discard("Mates")
    if "Posa mates" in found:
        found.discard("Mates")
    return next(iter(found)) if len(found) == 1 else None


def normalize_phone(raw: object) -> tuple[str | None, str | None]:
    value = clean(raw)
    if not value:
        return None, None
    digits = re.sub(r"\D", "", value)
    candidates = [value]
    if digits.startswith("00"):
        candidates.insert(0, "+" + digits[2:])
        if digits.startswith("0054") and not digits.startswith("00549"):
            candidates.insert(0, "+549" + digits[4:])
    elif not value.startswith("+") and len(digits) >= 10:
        candidates.insert(0, "+" + digits)
        if digits.startswith("54") and not digits.startswith("549"):
            candidates.insert(0, "+549" + digits[2:])
        if digits.startswith("598598"):
            candidates.insert(0, "+" + digits[3:])
    for candidate in dict.fromkeys(candidates):
        try:
            number = phonenumbers.parse(candidate, "UY")
        except phonenumbers.NumberParseException:
            continue
        if not phonenumbers.is_valid_number(number):
            continue
        if phonenumbers.number_type(number) not in (PhoneNumberType.MOBILE, PhoneNumberType.FIXED_LINE_OR_MOBILE):
            continue
        region = phonenumbers.region_code_for_number(number)
        return phonenumbers.format_number(number, PhoneNumberFormat.E164), COUNTRIES.get(region, region)
    return None, None


def prepare(source: Path) -> tuple[list[dict], list[dict], dict]:
    book = load_workbook(source, read_only=True, data_only=True)
    sheet = book["Contactos"]
    records: list[dict] = []
    review: list[dict] = []
    nonempty = 0
    sheet_rows = 0
    for row_number, cells in enumerate(sheet.iter_rows(min_row=6, max_col=6, values_only=True), start=6):
        sheet_rows += 1
        name = clean(cells[0])
        raw_phones = [clean(cell) for cell in cells[1:5] if clean(cell)]
        email_raw = clean(cells[5]).lower()
        if not name and not raw_phones and not email_raw:
            continue
        nonempty += 1
        issues: list[str] = []
        phones: list[str] = []
        countries: dict[str, str | None] = {}
        for raw in raw_phones:
            number, country = normalize_phone(raw)
            if number:
                if number not in phones:
                    phones.append(number)
                    countries[number] = country
            else:
                issues.append(f"Teléfono no válido: {raw}")
        email = email_raw if EMAIL.fullmatch(email_raw) else None
        if email_raw and not email:
            issues.append("Correo no válido")
        if not name or len(name) > 120:
            issues.append("Falta nombre o supera 120 caracteres")
        if not phones:
            issues.append("Falta celular válido")
        row = {"source_row": row_number, "name": name, "phones": phones, "countries": countries,
               "email": email, "raw_phones": raw_phones, "raw_email": email_raw}
        if not name or len(name) > 120 or not phones:
            review.append({**row, "reasons": issues})
            continue
        if issues:
            review.append({**row, "reasons": issues, "imported_with_warning": True})
        records.append(row)
    book.close()

    parent = list(range(len(records)))

    def root(i: int) -> int:
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    phone_owner: dict[str, int] = {}
    email_owner: dict[str, int] = {}
    for i, record in enumerate(records):
        for phone in record["phones"]:
            if phone in phone_owner:
                parent[root(i)] = root(phone_owner[phone])
            else:
                phone_owner[phone] = i
        if record["email"]:
            email = record["email"]
            if email in email_owner:
                parent[root(i)] = root(email_owner[email])
            else:
                email_owner[email] = i
    groups: dict[int, list[dict]] = defaultdict(list)
    for i, record in enumerate(records):
        groups[root(i)].append(record)

    imports: list[dict] = []
    duplicate_rows = 0
    for group in groups.values():
        names = {fold(item["name"]) for item in group}
        emails = {item["email"] for item in group if item["email"]}
        if len(names) > 1 or len(emails) > 1:
            reason = "Nombres distintos para el mismo teléfono/correo" if len(names) > 1 else "Correos distintos para el mismo contacto"
            review.extend({**item, "reasons": [reason]} for item in group)
            continue
        phones = list(dict.fromkeys(phone for item in group for phone in item["phones"]))
        country_set = {item["countries"].get(phone) for item in group for phone in item["phones"]}
        if len(country_set) > 1:
            review.extend({**item, "reasons": ["Teléfonos de países diferentes"]} for item in group)
            continue
        first = group[0]
        primary = phones[0]
        imports.append({
            "agenda_import_key": hashlib.sha256(primary.encode("ascii")).hexdigest(),
            "full_name": first["name"], "phone": primary, "email": next(iter(emails), None),
            "country": first["countries"].get(primary), "agenda_additional_phones": phones[1:],
            "agenda_inferred_category": category_for(first["name"]),
            "source_rows": [item["source_row"] for item in group],
        })
        duplicate_rows += len(group) - 1
    imports.sort(key=lambda item: item["agenda_import_key"])
    summary = {"sheet_rows": sheet_rows, "nonempty_rows": nonempty, "eligible_rows": len(records),
               "unique_imports": len(imports), "merged_duplicate_rows": duplicate_rows,
               "review_rows": len({item["source_row"] for item in review}),
               "inferred_categories": dict(Counter(item["agenda_inferred_category"] for item in imports if item["agenda_inferred_category"]))}
    return imports, review, summary


def reconcile_existing(imports: list[dict], review: list[dict], existing: list[dict]) -> tuple[list[dict], list[dict]]:
    by_phone: dict[str, set[str]] = defaultdict(set)
    by_email: dict[str, set[str]] = defaultdict(set)
    by_id = {row["id"]: row for row in existing}
    for row in existing:
        phone, _ = normalize_phone(row.get("phone"))
        if phone:
            by_phone[phone].add(row["id"])
        email = clean(row.get("email")).lower()
        if email:
            by_email[email].add(row["id"])
    accepted = []
    for item in imports:
        matches: set[str] = set()
        for phone in [item["phone"], *item["agenda_additional_phones"]]:
            matches.update(by_phone[phone])
        if item["email"]:
            matches.update(by_email[item["email"]])
        if len(matches) > 1:
            review.append({**item, "reasons": ["Coincide con más de una ficha existente"]})
            continue
        if matches:
            existing_row = by_id[next(iter(matches))]
            old_email = clean(existing_row.get("email")).lower()
            if fold(existing_row["full_name"]) != fold(item["full_name"]) or (old_email and item["email"] and old_email != item["email"]):
                review.append({**item, "reasons": ["Conflicto de nombre o correo con ficha existente"], "existing_id": existing_row["id"]})
                continue
            item["existing_id"] = existing_row["id"]
            old_phone, _ = normalize_phone(existing_row.get("phone"))
            if old_phone and old_phone not in [item["phone"], *item["agenda_additional_phones"]]:
                item["agenda_additional_phones"].append(old_phone)
        accepted.append(item)
    return accepted, review


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--review-output", type=Path, help="Copia duradera del informe de revisión")
    parser.add_argument("--existing-json", type=Path, help="Fichas minoristas actuales para detectar conflictos antes de cargar")
    args = parser.parse_args()
    imports, review, summary = prepare(args.source)
    if args.existing_json:
        existing = json.loads(args.existing_json.read_text(encoding="utf-8"))
        imports, review = reconcile_existing(imports, review, existing)
        summary["unique_imports"] = len(imports)
        summary["matched_existing"] = sum("existing_id" in item for item in imports)
    accepted_rows = {number for item in imports for number in item["source_rows"]}
    reviewed_rows = {number for item in review for number in (item.get("source_rows") or [item.get("source_row")])}
    summary["merged_duplicate_rows"] = sum(len(item["source_rows"]) - 1 for item in imports)
    summary["accepted_source_rows"] = len(accepted_rows)
    summary["review_rows"] = len(reviewed_rows)
    summary["quarantined_rows"] = len(reviewed_rows - accepted_rows)
    summary["imported_with_warning_rows"] = len(reviewed_rows & accepted_rows)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    (args.output_dir / "agenda_import_rows.json").write_text(json.dumps(imports, ensure_ascii=False, indent=2), encoding="utf-8")
    review_json = json.dumps({"summary": summary, "rows": review}, ensure_ascii=False, indent=2)
    (args.output_dir / "agenda_import_review.json").write_text(review_json, encoding="utf-8")
    if args.review_output:
        args.review_output.parent.mkdir(parents=True, exist_ok=True)
        args.review_output.write_text(review_json, encoding="utf-8")
    def sql_literal(rows: list[dict]) -> str:
        return "'" + json.dumps(rows, ensure_ascii=False, separators=(",", ":")).replace("'", "''") + "'::jsonb"

    existing_rows = [{key: value for key, value in item.items() if key != "source_rows"} for item in imports if item.get("existing_id")]
    if existing_rows:
        existing_sql = f"""with source as (
  select * from jsonb_to_recordset({sql_literal(existing_rows)}) as x(
    existing_id uuid, agenda_import_key text, phone text, email text, country text,
    agenda_additional_phones text[], agenda_inferred_category text)
)
update public.local_sales_customers c
set agenda_import_key = source.agenda_import_key,
    phone = source.phone, email = coalesce(c.email, source.email),
    country = coalesce(c.country, source.country),
    agenda_additional_phones = source.agenda_additional_phones,
    agenda_inferred_category = source.agenda_inferred_category
from source where c.id = source.existing_id and c.agenda_import_key is null;
"""
        (args.output_dir / "agenda_existing_updates.sql").write_text(existing_sql, encoding="utf-8")
    for offset in range(0, len(imports), 150):
        rows = [{key: value for key, value in item.items() if key != "source_rows"} for item in imports[offset:offset + 150]]
        (args.output_dir / f"agenda_batch_{offset // 150:03d}.json").write_text(json.dumps(rows, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
        new_rows = [item for item in rows if not item.get("existing_id")]
        if new_rows:
            batch_sql = f"""with source as (
  select * from jsonb_to_recordset({sql_literal(new_rows)}) as x(
    agenda_import_key text, full_name text, phone text, email text, country text,
    agenda_additional_phones text[], agenda_inferred_category text)
)
insert into public.local_sales_customers
  (agenda_import_key, full_name, phone, email, country, agenda_additional_phones, agenda_inferred_category)
select agenda_import_key, full_name, phone, email, country, agenda_additional_phones, agenda_inferred_category
from source
on conflict (agenda_import_key) where agenda_import_key is not null do nothing;
"""
            (args.output_dir / f"agenda_batch_{offset // 150:03d}.sql").write_text(batch_sql, encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False))


if __name__ == "__main__":
    main()
